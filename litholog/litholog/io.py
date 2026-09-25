"""Read projects from Excel workbooks or CSV folders, and write the input template.

Headers are matched loosely: case, spaces, units in brackets and common
synonyms are ignored, so "Borehole ID", "BH_ID" and "Well" all work.
"""

from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd

from .patterns import DEFAULT_LEGEND, PATTERNS, Legend
from .project import TABLE_COLUMNS, Project, empty_table

SHEET_ALIASES = {
    "boreholes": ["boreholes", "borehole", "collars", "collar", "locations", "wells", "sites"],
    "lithology": ["lithology", "litho", "strata", "lithologs", "litholog", "geology"],
    "construction": ["construction", "well_construction", "casing", "completion"],
    "water_levels": ["water_levels", "water_level", "waterlevels", "wl", "swl"],
    "downhole": ["downhole", "downhole_data", "geophysics", "logs", "measurements"],
    "legend": ["legend", "lithology_legend", "codes", "lithology_codes"],
}

COLUMN_ALIASES = {
    "boreholes": {
        "borehole_id": ["borehole_id", "borehole", "bh_id", "bh", "well_id", "well", "id", "hole_id", "name"],
        "x": ["x", "easting", "longitude", "lon", "long"],
        "y": ["y", "northing", "latitude", "lat"],
        "elevation": ["elevation", "elev", "rl", "ground_level", "ground_elevation", "gl", "z"],
        "total_depth": ["total_depth", "td", "depth", "drilled_depth", "total_drilled_depth"],
    },
    "lithology": {
        "borehole_id": ["borehole_id", "borehole", "bh_id", "bh", "well_id", "well", "id", "hole_id"],
        "from": ["from", "depth_from", "top", "from_depth", "top_depth"],
        "to": ["to", "depth_to", "bottom", "to_depth", "bottom_depth", "base"],
        "code": ["code", "lithology_code", "litho_code", "lith_code", "rock_code", "lithology"],
        "description": ["description", "desc", "remarks", "log_description"],
    },
    "construction": {
        "borehole_id": ["borehole_id", "borehole", "bh_id", "bh", "well_id", "well", "id", "hole_id"],
        "from": ["from", "depth_from", "top"],
        "to": ["to", "depth_to", "bottom"],
        "element": ["element", "type", "component", "item"],
        "diameter": ["diameter", "dia", "size"],
        "material": ["material", "remarks"],
    },
    "water_levels": {
        "borehole_id": ["borehole_id", "borehole", "bh_id", "bh", "well_id", "well", "id", "hole_id"],
        "date": ["date", "measured_on", "date_measured"],
        "depth": ["depth", "depth_to_water", "dtw", "swl", "water_level", "wl"],
    },
    "downhole": {
        "borehole_id": ["borehole_id", "borehole", "bh_id", "bh", "well_id", "well", "id", "hole_id"],
        "depth": ["depth", "md"],
        "parameter": ["parameter", "param", "variable", "log", "curve"],
        "value": ["value", "reading", "val"],
        "unit": ["unit", "units"],
    },
}

NUMERIC = {
    "boreholes": ["x", "y", "elevation", "total_depth"],
    "lithology": ["from", "to"],
    "construction": ["from", "to", "diameter"],
    "water_levels": ["depth"],
    "downhole": ["depth", "value"],
}


def normalize(name) -> str:
    s = re.sub(r"[\(\[].*?[\)\]]", "", str(name)).strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def _match_table(sheet_name: str):
    n = normalize(sheet_name)
    for table, aliases in SHEET_ALIASES.items():
        if n in aliases:
            return table
    return None


def _standardize(df: pd.DataFrame, table: str) -> pd.DataFrame:
    df = df.dropna(how="all").copy()
    aliases = COLUMN_ALIASES.get(table, {})
    rename, used = {}, set()
    norm = {c: normalize(c) for c in df.columns}
    # Exact canonical names first, then synonyms in priority order.
    for canon, alts in aliases.items():
        for alt in alts:
            hit = next((c for c, n in norm.items() if n == alt and c not in used), None)
            if hit is not None:
                rename[hit] = canon
                used.add(hit)
                break
    for c, n in norm.items():
        if c not in used:
            rename[c] = n
    df = df.rename(columns=rename)
    df = df.loc[:, ~df.columns.duplicated()]
    for col in TABLE_COLUMNS.get(table, []):
        if col not in df.columns:
            df[col] = np.nan
    if "borehole_id" in df.columns:
        df = df[df["borehole_id"].notna()]
        df["borehole_id"] = df["borehole_id"].map(_clean_id)
    for col in NUMERIC.get(table, []):
        df[col] = pd.to_numeric(df[col], errors="coerce")
    if table == "lithology":
        df["code"] = df["code"].map(lambda v: "" if pd.isna(v) else str(v).strip().upper())
    if table == "construction":
        df["element"] = df["element"].map(canonical_element)
    if table == "water_levels":
        df["date"] = df["date"].map(parse_date)
    return df.reset_index(drop=True)


def parse_date(v):
    """ISO dates (2025-03-13) as-is; otherwise Indian/European day-first (13-03-2025)."""
    if pd.isna(v) or str(v).strip() == "":
        return pd.NaT
    if isinstance(v, (pd.Timestamp, np.datetime64)) or hasattr(v, "year"):
        return pd.Timestamp(v)
    text = str(v).strip()
    iso = re.match(r"^\d{4}[-/]", text) is not None
    return pd.to_datetime(text, errors="coerce", dayfirst=not iso, yearfirst=iso)


def _clean_id(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


ELEMENTS = {
    "casing": ["casing", "plain_casing", "blank_casing", "blank", "plain", "pipe", "riser"],
    "screen": ["screen", "slotted", "slotted_pipe", "slotted_casing", "filter", "strainer"],
    "gravel_pack": ["gravel_pack", "gravel", "filter_pack", "pack", "shrouding"],
    "seal": ["seal", "grout", "cement", "bentonite", "clay_seal", "sanitary_seal"],
    "open_hole": ["open_hole", "open", "openhole", "uncased"],
}


def canonical_element(v) -> str:
    n = normalize(v) if pd.notna(v) else ""
    for canon, alts in ELEMENTS.items():
        if n in alts:
            return canon
    return n or "casing"


def load_project(path, name: str | None = None) -> Project:
    """Load an Excel workbook (.xlsx/.xls) or a folder of CSV files."""
    path = Path(path)
    raw: dict[str, pd.DataFrame] = {}
    if path.is_dir():
        for f in sorted(path.glob("*.csv")):
            table = _match_table(f.stem)
            if table:
                raw[table] = pd.read_csv(f)
    elif path.suffix.lower() == ".csv":
        raise ValueError("Give the folder that holds the CSV files, not a single CSV")
    else:
        for sheet, df in pd.read_excel(path, sheet_name=None).items():
            table = _match_table(sheet)
            if table and table not in raw:
                raw[table] = df
    if "boreholes" not in raw and "lithology" not in raw:
        raise ValueError(
            f"No 'Boreholes' or 'Lithology' sheet found in {path}. "
            "Run `litholog template` to get a correctly laid-out workbook."
        )

    tables = {t: _standardize(raw[t], t) if t in raw else empty_table(t) for t in TABLE_COLUMNS}

    # Boreholes present only in data tables still get a (location-less) collar row.
    known = set(tables["boreholes"]["borehole_id"])
    extra = []
    for t in ("lithology", "construction", "water_levels", "downhole"):
        for bid in tables[t]["borehole_id"]:
            if bid not in known:
                known.add(bid)
                extra.append({"borehole_id": bid})
    if extra:
        tables["boreholes"] = pd.concat([tables["boreholes"], pd.DataFrame(extra)], ignore_index=True)

    legend = Legend()
    if "legend" in raw:
        leg = raw["legend"].dropna(how="all").copy()
        leg.columns = [normalize(c) for c in leg.columns]
        leg = leg.rename(columns={"colour": "color", "fill": "color", "lithology": "name"})
        legend = legend.updated(leg.to_dict("records"))

    return Project(**tables, legend=legend, name=name or path.stem)


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------

_INSTRUCTIONS = [
    ("LithoLog input workbook", ""),
    ("", ""),
    ("How to use", "Fill in the sheets below (keep the sheet names), save, then run:"),
    ("", "litholog striplog this_file.xlsx"),
    ("", ""),
    ("Boreholes", "One row per borehole. Borehole ID is required; any extra columns you add "
                  "(Location, Date, Drilling method, Logged by, ...) are printed in the log header."),
    ("Lithology", "One row per interval: Borehole ID, From (m), To (m), Code, Description. "
                  "Codes come from the Legend sheet; add your own codes there."),
    ("Construction", "Optional. Element = casing, screen, gravel pack, seal (grout) or open hole; "
                     "Diameter in mm."),
    ("WaterLevels", "Optional. Depth to water in m below ground level, with date."),
    ("Downhole", "Optional. Long format: one row per reading (Depth, Parameter, Value, Unit), "
                 "e.g. resistivity, gamma, EC, yield."),
    ("Legend", "Code, Name, Color (hex like #F3E196), Pattern. Patterns can be combined with '+', "
               "e.g. crosses+fractures. Available patterns: " + ", ".join(PATTERNS)),
    ("", ""),
    ("Depths", "All depths are metres below ground level (m bgl). Elevation in m above mean sea level."),
]

_EXAMPLE = {
    "Boreholes": pd.DataFrame([
        {"Borehole ID": "BH-01", "Easting": 650120.0, "Northing": 1689340.0, "Elevation (m amsl)": 482.5,
         "Total depth (m)": 60.0, "Location": "Example site", "Date": "2025-03-12",
         "Drilling method": "DTH", "Logged by": ""},
        {"Borehole ID": "BH-02", "Easting": 650410.0, "Northing": 1689105.0, "Elevation (m amsl)": 479.0,
         "Total depth (m)": 45.0, "Location": "Example site", "Date": "2025-03-14",
         "Drilling method": "DTH", "Logged by": ""},
    ]),
    "Lithology": pd.DataFrame([
        ("BH-01", 0, 1.5, "RSOIL", "Red sandy soil"),
        ("BH-01", 1.5, 9, "WGRA", "Highly weathered pink granite"),
        ("BH-01", 9, 14, "FGRA", "Fractured granite, water struck at 12 m"),
        ("BH-01", 14, 60, "GRA", "Massive grey granite"),
        ("BH-02", 0, 2, "RSOIL", "Red soil"),
        ("BH-02", 2, 11, "WGN", "Weathered gneiss"),
        ("BH-02", 11, 45, "GN", "Hard gneiss"),
    ], columns=["Borehole ID", "From (m)", "To (m)", "Code", "Description"]),
    "Construction": pd.DataFrame([
        ("BH-01", 0, 15, "casing", 165, "PVC"),
        ("BH-01", 15, 60, "open hole", 150, ""),
        ("BH-02", 0, 12, "casing", 165, "MS"),
        ("BH-02", 12, 45, "open hole", 150, ""),
    ], columns=["Borehole ID", "From (m)", "To (m)", "Element", "Diameter (mm)", "Material"]),
    "WaterLevels": pd.DataFrame([
        ("BH-01", "2025-03-13", 11.2),
        ("BH-02", "2025-03-15", 9.8),
    ], columns=["Borehole ID", "Date", "Depth to water (m bgl)"]),
    "Downhole": pd.DataFrame(columns=["Borehole ID", "Depth (m)", "Parameter", "Value", "Unit"]),
}


def write_template(path, example: bool = True) -> Path:
    """Write an empty (or example-filled) input workbook."""
    path = Path(path)
    sheets = {k: (v if example else v.iloc[0:0]) for k, v in _EXAMPLE.items()}
    legend = pd.DataFrame([
        {"Code": t.code, "Name": t.name, "Color": t.color, "Pattern": t.pattern, "Group": t.group}
        for t in DEFAULT_LEGEND.values()
    ])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(_INSTRUCTIONS, columns=["Topic", "Notes"]).to_excel(xw, sheet_name="Instructions", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        legend.to_excel(xw, sheet_name="Legend", index=False)
        _style_workbook(xw.book)
    return path


def _style_workbook(book):
    from openpyxl.styles import Font, PatternFill

    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for ws in book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 90)
    legend = book["Legend"]
    for row in legend.iter_rows(min_row=2):
        color = str(row[2].value or "").lstrip("#")
        if re.fullmatch(r"[0-9A-Fa-f]{6}", color):
            row[2].fill = PatternFill("solid", fgColor=color.upper())
